from pathlib import Path

from docx import Document

from backend.app.ingestion.base import ParsedDocument, TextBlock
from backend.app.ingestion.parsers.base_parser import BaseParser


class DocxParser(BaseParser):
    extensions = frozenset({".docx"})

    def parse(self, path: Path) -> ParsedDocument:
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        return ParsedDocument(blocks=[TextBlock(page_number=1, text=full_text)])


def parse_docx(path: Path) -> ParsedDocument:
    return DocxParser().parse(path)
